from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os
import re

os.chdir('/sessions/great-fervent-euler/mnt/outputs')

# Mapping de archivos a semanas
semanas = {
    9: ("16-20 Marzo", "2-6 Marzo"),
    10: ("24-28 Marzo", "9-13 Marzo"),
    11: ("6-10 Abril", "16-20 Marzo"),
    12: ("13-17 Abril", "24-28 Marzo"),
    13: ("20-24 Abril", "6-10 Abril"),
    14: ("27-30 Abril", "13-17 Abril"),
    15: ("4-8 Mayo", "20-24 Abril"),
    16: ("11-15 Mayo", "27-30 Abril"),
    17: ("18-22 Mayo", "4-8 Mayo"),
    18: ("25-29 Mayo", "11-15 Mayo"),
}

# Areas mapping a tablas
areas_config = {
    'Apoyo Logístico': 0,
    'Gestión Sistemas': 1,
    'Soporte Telemático': 2,
    'Soporte INGENI@': 2,
    'Documental CENDOI': 3,
    'Gestión Proyectos': 4,
    'INGENI@': 5,
    'Producción': 6,
    'Administrativa': 7,
}

# Crear workbook
wb = Workbook()
ws = wb.active
ws.title = "Consolidado"

# Estilos
header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF", size=11)
border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Headers
headers = ["Área"] + [f"Sem {i+1}\n({semanas[i][0][:5]})" for i in range(9, 19)]
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col)
    cell.value = header
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = border

# Extender datos de archivos
print("="*80)
print("CONSOLIDANDO DATOS EN EXCEL")
print("="*80)

datos_consolidados = {}

for num_archivo in range(9, 19):
    filename = f"{num_archivo:02d}- Informe Semanal de Actividades.docx"
    
    if not os.path.exists(filename):
        print(f"  ✗ {filename} no encontrado")
        continue
    
    print(f"  ✓ Leyendo {filename}")
    
    doc = Document(filename)
    
    # Tabla 1: Apoyo Logístico
    if len(doc.tables) > 0:
        table = doc.tables[0]
        for row_idx, row in enumerate(table.rows[1:], 1):
            if len(row.cells) > 1:
                actividad = row.cells[0].text.strip()
                observacion = row.cells[1].text
                
                # Extraer el primer número significativo
                numeros = [int(x) for x in re.findall(r'\b(\d+)\b', observacion) if int(x) < 10000]
                
                if numeros and actividad not in ['Infraestructura']:
                    key = f"Apoyo Logístico - {actividad}"
                    if key not in datos_consolidados:
                        datos_consolidados[key] = {}
                    datos_consolidados[key][num_archivo] = numeros[0]
                elif actividad == 'Videoconferencia' and numeros:
                    # Para videoconferencia, tomar el número de sesiones
                    key = "Apoyo Logístico - Videoconferencias"
                    if key not in datos_consolidados:
                        datos_consolidados[key] = {}
                    datos_consolidados[key][num_archivo] = numeros[0]

# Escribir datos
row_num = 2
for area_name in sorted(datos_consolidados.keys()):
    cell_area = ws.cell(row=row_num, column=1)
    cell_area.value = area_name[:30]
    cell_area.border = border
    
    data_row = datos_consolidados[area_name]
    for col, num_archivo in enumerate(range(9, 19), 2):
        cell = ws.cell(row=row_num, column=col)
        if num_archivo in data_row:
            cell.value = data_row[num_archivo]
        cell.alignment = Alignment(horizontal='center')
        cell.border = border
    
    row_num += 1

# Ajustar ancho de columnas
ws.column_dimensions['A'].width = 35
for col in range(2, 12):
    ws.column_dimensions[chr(64 + col)].width = 12

# Guardar
output_file = '/sessions/great-fervent-euler/mnt/outputs/DRAI_Consolidado_09-18.xlsx'
wb.save(output_file)

print(f"\n✓ Archivo guardado: DRAI_Consolidado_09-18.xlsx")
print(f"  Áreas procesadas: {len(datos_consolidados)}")
print(f"  Semanas: 10 (archivos 09-18)")

