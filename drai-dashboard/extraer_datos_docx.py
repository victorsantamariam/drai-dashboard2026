from docx import Document
import os
import json
from pathlib import Path

# Cambiar al directorio de outputs
os.chdir('/sessions/great-fervent-euler/mnt/outputs')

# Diccionario para almacenar datos consolidados
datos_consolidados = {
    'Apoyo Logístico y Videoconferencia': [],
    'Gestión de Sistemas de Información': [],
    'Soporte Telemático': [],
    'Soporte Técnico y Académico INGENI@': [],
    'Gestión Documental CENDOI': [],
    'Unidad de Gestión de Proyectos': [],
    'INGENI@': [],
    'Producción': [],
    'Gestión Administrativa': []
}

# Procesar archivos 09-18
print("="*80)
print("EXTRAYENDO DATOS DE ARCHIVOS 09-18")
print("="*80)

for num_archivo in range(9, 19):
    filename = f"{num_archivo:02d}- Informe Semanal de Actividades.docx"
    
    if not os.path.exists(filename):
        print(f"\n✗ No encontrado: {filename}")
        continue
    
    print(f"\n{'─'*80}")
    print(f"PROCESANDO: {filename}")
    print(f"{'─'*80}")
    
    try:
        doc = Document(filename)
        
        print(f"  Tablas encontradas: {len(doc.tables)}")
        
        # Extraer todas las tablas
        for table_idx, table in enumerate(doc.tables):
            rows = table.rows
            print(f"  Tabla {table_idx + 1}: {len(rows)} filas × {len(rows[0].cells) if rows else 0} columnas")
            
            # Convertir tabla a datos
            table_data = []
            for row in rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_data.append(row_data)
            
            # Buscar la tabla principal con actividades
            if len(table_data) > 2 and len(table_data[0]) >= 2:
                # Primera fila es header, siguiente filas son datos
                header = table_data[0]
                
                # Si parece tabla de actividades (columnas: Actividad/Proyecto, Observación)
                if 'actividad' in ' '.join(header).lower() or 'proyecto' in ' '.join(header).lower():
                    print(f"\n    ✓ Tabla de actividades detectada (Tabla {table_idx + 1}):")
                    
                    # Procesar filas de datos
                    for row_idx in range(1, len(table_data)):
                        row = table_data[row_idx]
                        if len(row) >= 1 and row[0].strip():  # Si la primera celda no está vacía
                            actividad = row[0].strip()
                            print(f"      - {actividad[:60]}")
        
        print(f"\n  ✓ Archivo {num_archivo:02d} procesado exitosamente")
        
    except Exception as e:
        print(f"  ✗ Error procesando {filename}: {e}")

print("\n" + "="*80)
print("EXTRACCIÓN COMPLETADA")
print("="*80)

