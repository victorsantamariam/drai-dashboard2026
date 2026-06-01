from docx import Document
import os
import re

os.chdir('/sessions/great-fervent-euler/mnt/outputs')

# Semanas correspondientes a cada archivo
semanas_mapping = {
    9: "16-20 Marzo",
    10: "24-28 Marzo",
    11: "6-10 Abril",
    12: "13-17 Abril",
    13: "20-24 Abril",
    14: "27-30 Abril",
    15: "4-8 Mayo",
    16: "11-15 Mayo",
    17: "18-22 Mayo",
    18: "25-29 Mayo"
}

# Estructura para almacenar datos
datos = {}

print("="*80)
print("EXTRAYENDO NÚMEROS CONSOLIDADOS DE ARCHIVOS 09-18")
print("="*80)

for num_archivo in range(9, 19):
    filename = f"{num_archivo:02d}- Informe Semanal de Actividades.docx"
    semana = semanas_mapping.get(num_archivo, f"Semana {num_archivo}")
    
    if not os.path.exists(filename):
        continue
    
    print(f"\n{'─'*80}")
    print(f"ARCHIVO {num_archivo:02d} - Semana {semana}")
    print(f"{'─'*80}")
    
    doc = Document(filename)
    
    # Tabla 1 es siempre Apoyo Logístico
    if len(doc.tables) > 0:
        table = doc.tables[0]
        print("\n  ÁREA 1: APOYO LOGÍSTICO Y VIDEOCONFERENCIA")
        
        # Procesar filas (header en índice 0, datos desde índice 1)
        for row_idx, row in enumerate(table.rows[1:], 1):
            actividad = row.cells[0].text.strip()
            observacion = row.cells[1].text if len(row.cells) > 1 else ""
            
            # Extraer números
            numeros = re.findall(r'\b(\d+)\b', observacion)
            if numeros:
                print(f"    {actividad[:30]:30s} | Números: {numeros[:5]}")
    
    # Tabla 2 - Gestión de Sistemas (varía con índices)
    if len(doc.tables) > 1:
        print(f"\n  ÁREA 2: GESTIÓN DE SISTEMAS")
        table = doc.tables[1]
        for row_idx, row in enumerate(table.rows[1:3], 1):  # Solo primeras 2
            actividad = row.cells[0].text.strip()
            observacion = row.cells[1].text if len(row.cells) > 1 else ""
            numeros = re.findall(r'\b(\d+)\b', observacion)
            if numeros:
                print(f"    {actividad[:30]:30s} | Números: {numeros[:3]}")
    
    # Tabla 3 - Soporte Telemático/Académico
    if len(doc.tables) > 2:
        print(f"\n  ÁREA 3/4: SOPORTE TELEMÁTICO Y ACADÉMICO")
        table = doc.tables[2]
        for row_idx, row in enumerate(table.rows[1:4], 1):  # Primeras 3
            actividad = row.cells[0].text.strip()
            observacion = row.cells[1].text if len(row.cells) > 1 else ""
            numeros = re.findall(r'\b(\d+)\b', observacion)
            if numeros:
                print(f"    {actividad[:30]:30s} | Números: {numeros[:3]}")

print("\n" + "="*80)
print("EXTRACCIÓN COMPLETADA")
print("="*80)

