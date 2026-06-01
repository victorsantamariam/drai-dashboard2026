from flask import Flask, request, jsonify
from flask_cors import CORS
from docx import Document
import os
import zipfile
import io
import re
from datetime import datetime
import json

app = Flask(__name__)
CORS(app)

database = {'reports': [], 'files': []}

def extract_numbers_from_text(text):
    numbers = re.findall(r'\b(\d+)\b', text)
    return [int(n) for n in numbers if int(n) < 100000]

def process_docx_file(file_path):
    try:
        doc = Document(file_path)
        data = {'tables': [], 'paragraphs': [], 'numbers': []}
        for table_idx, table in enumerate(doc.tables):
            table_data = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_data.append(row_data)
            if table_data:
                data['tables'].append({'index': table_idx, 'rows': table_data})
        for para in doc.paragraphs:
            if para.text.strip():
                data['paragraphs'].append(para.text.strip())
                numbers = extract_numbers_from_text(para.text)
                data['numbers'].extend(numbers)
        return data
    except Exception as e:
        return {'error': str(e)}

def consolidate_data(files_data):
    consolidated = {
        'total_files': len(files_data),
        'total_tables': sum(len(f['data']['tables']) for f in files_data),
        'total_activities': 0,
        'files_processed': []
    }
    for file_info in files_data:
        file_name = file_info['filename']
        numbers = file_info['data'].get('numbers', [])
        if numbers:
            total = sum(numbers[:10])
            consolidated['total_activities'] += total
            consolidated['files_processed'].append({'name': file_name, 'activities': total})
    return consolidated

@app.route('/api/health', methods=['GET'])
def health():
    return jsonify({'status': 'ok', 'message': 'Backend running'}), 200

@app.route('/api/upload', methods=['POST'])
def upload_files():
    try:
        if 'files' not in request.files:
            return jsonify({'error': 'No files provided'}), 400
        files = request.files.getlist('files')
        results = []
        all_files_data = []
        for file in files:
            if file.filename == '':
                continue
            filename = file.filename
            if filename.endswith('.zip'):
                with zipfile.ZipFile(io.BytesIO(file.read())) as zip_ref:
                    for zip_file in zip_ref.namelist():
                        if zip_file.endswith('.docx'):
                            with zip_ref.open(zip_file) as docx_file:
                                doc_data = process_docx_file(io.BytesIO(docx_file.read()))
                                all_files_data.append({'filename': zip_file, 'data': doc_data})
                                results.append({'filename': zip_file, 'status': 'processed'})
            elif filename.endswith('.docx'):
                doc_data = process_docx_file(file)
                all_files_data.append({'filename': filename, 'data': doc_data})
                results.append({'filename': filename, 'status': 'processed'})
        consolidated = consolidate_data(all_files_data)
        report = {'id': len(database['reports']) + 1, 'timestamp': datetime.now().isoformat(), 'consolidated': consolidated, 'files': results}
        database['reports'].append(report)
        return jsonify({'success': True, 'report_id': report['id'], 'data': consolidated, 'files': results}), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/reports', methods=['GET'])
def get_reports():
    return jsonify({'reports': database['reports']}), 200

@app.route('/api/reports/<int:report_id>', methods=['GET'])
def get_report(report_id):
    for report in database['reports']:
        if report['id'] == report_id:
            return jsonify(report), 200
    return jsonify({'error': 'Report not found'}), 404

@app.route('/api/reports/<int:report_id>', methods=['DELETE'])
def delete_report(report_id):
    for i, report in enumerate(database['reports']):
        if report['id'] == report_id:
            database['reports'].pop(i)
            return jsonify({'success': True}), 200
    return jsonify({'error': 'Report not found'}), 404

@app.route('/api/export/<int:report_id>', methods=['GET'])
def export_report(report_id):
    for report in database['reports']:
        if report['id'] == report_id:
            return jsonify(report), 200
    return jsonify({'error': 'Report not found'}), 404

@app.route('/', methods=['GET'])
def index():
    return jsonify({'message': 'DRAI Dashboard API', 'version': '1.0'}), 200

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
